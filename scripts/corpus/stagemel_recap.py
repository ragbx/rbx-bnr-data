#!/usr/bin/env python3
r"""
stagemel_recap.py — Recap Excel par corpus, à partir de cas1/cas2/cas3.
-> results/corpus/stagemel/<date>/<corpus>_recap_<date>.xlsx (colonnes à la largeur du texte, filtres sur les en-têtes, en-tête figé)

Reproduit la même structure que les recap_<code>_<date>.xlsx de la stagiaire
(CAS / <CODE> / STATUT / CHEMIN / PROBLEMES / À FAIRE), mais seulement pour ce
qui est mécaniquement dérivable :

  - UUID : ajouté quand disponible — celui du REF pour cas1/cas3 ; pour cas2,
           uniquement si un candidat ERREUR DAO a été rapproché (uuid du cas1/cas3
           correspondant), sinon vide (par définition, un vrai cas2 n'a pas d'uuid).

  - CAS : renommé en libellés courts plutôt que « CAS 1/2/3 » — APPARIÉ (cas1,
          dans REF+DAO), DAO SEUL (cas2), REF SEUL (cas3).

  - cas1 (APPARIÉ) : STATUT = conservation_statut du REF, tel quel ;
           CHEMIN = path du REF ;
           PROBLEMES = « Pas de format tif » si le fichier est un .jpg/.jpeg sans
           .tif/.tiff de même clé parmi les fichiers REF du corpus (cas1+cas3,
           casse des extensions et des clés ignorée), sinon vide ;
           À FAIRE = déduit du statut via ACTION_SUR (statuts harmonisés du ref,
           qui sont déjà des décisions). La famille INCONNU* en est absente :
           cellule vide et avertissement, la politique à appliquer (cf. CONFIG
           de med_s3_key_cible.py) doit être validée, elle n'est pas inventée ici.
           Exclues du recap : les lignes À SUPPRIMER (*) dont le master (.tif)
           est déjà TRANSFERT_S3_OK — recherché sur tout le ref, tous corpus
           confondus (pas seulement celui traité), rien à trancher dessus. Le
           ref consulté est celui de l'extraction (_ref_<date>.txt), à défaut
           le plus récent.

  - cas2 (DAO SEUL) : STATUT = « SEUL DAO » par défaut (DAO attribuée mais fichier
           introuvable dans REF) ; sauf si sa clé normalisée (casse/ponctuation
           ignorées) correspond à une clé du cas1/cas3, auquel cas STATUT =
           « ERREUR DAO (candidat) » — À VALIDER contre la notice EAD, cf.
           méthodologie du stage (les cas d'erreur de nommage réels doivent être
           confirmés à la main, ce script ne fait que proposer le rapprochement).
           À FAIRE toujours vide : ni « À numériser » pour les SEUL DAO (fichier
           introuvable ≠ document non numérisé, retiré le 2026-09-25), ni pour
           les candidats ERREUR DAO (dépend de la confirmation). La correspondance exacte
           cherche d'abord parmi les fichiers pas encore À SUPPRIMER, et ne se
           rabat sur un fichier À SUPPRIMER que si rien d'autre n'a matché
           (STATUT le signale alors explicitement) : pointer vers un fichier
           voué à disparaître serait trompeur si un autre candidat existe.
           Pas de rapprochement flou/approchant : un fichier « non retrouvé
           dans REF » reste tel quel, sans piste automatique proposée — trop
           de faux positifs sur les corpus à foliotation dense (cf. mémoire :
           tout folio voisin d'un manuscrit MED_MS ressortait à >0.85 sans lien
           réel), retiré le 2026-08-22.

  - cas3 (REF SEUL) : fichier du REF référencé par aucune notice EAD.
           STATUT = conservation_statut du REF ; CHEMIN = path du REF ;
           PROBLEMES = « Absent des notices EAD » (+ « Pas de format tif », même
           règle que cas1) ; À FAIRE volontairement vide (notice à créer ?
           fichier hors périmètre ? — décision de l'utilisateur). Même exclusion
           des À SUPPRIMER dont le master est déjà TRANSFERT_S3_OK.

  - Non couvert : les fichiers « Zébulon » (aucune source de données dans le
    dépôt pour ce dossier partagé) -- absent de ce recap.

Usage
-----
  conda run -n rbx-bnr-data python scripts/corpus/stagemel_recap.py MUS_ARC
  conda run -n rbx-bnr-data python scripts/corpus/stagemel_recap.py MUS_ARC --date 20260821
"""
import argparse
import re
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).parent))
from stagemel_extraction_corpus import dernier_ref, dossier_date, dossier_travail, ref_trace_path  # noqa: E402


# Actions sans ambiguïté : les statuts harmonisés du ref (fusion 2026-07-11) sont
# déjà des décisions prises, pas des diagnostics à interpréter — cf. mémoire
# project_ref_merge_20260711. Seule la famille INCONNU* reste une politique à
# définir par corpus (cf. CONFIG de scripts/s3/med_s3_key_cible.py), donc
# volontairement absente d'ACTION_SUR.
ACTION_SUR = {
    "TRANSFERT_S3_OK": "Rien à faire",
    "À TRANSFERER": "À transférer",
    "À TRANSFERER APRES VALIDATION": "À transférer (après validation)",
    "À TRANSFERER VOIR MARIE": "À transférer (voir Marie)",
    "EN LIGNE - À TRANSFERER ?": "À transférer (à confirmer)",
    "À SUPPRIMER (DIFFUSION)": "À supprimer",
    "À SUPPRIMER (REMPLACEMENT)": "À supprimer",
    "À SUPPRIMER (DOUBLON)": "À supprimer",
    "À SUPPRIMER (DOUBLON AZ)": "À supprimer",
    "À SUPPRIMER (DOUBLON S3-AZ)": "À supprimer",
    "À SUPPRIMER (FILE_TYPE)": "À supprimer",
    "À SUPPRIMER (RENOMMAGE)": "À supprimer",
    "À SUPPRIMER (MED_PAR)": "À supprimer",
    "À SUPPRIMER (TESTS)": "À supprimer",
    "À SUPPRIMER (DOUBLON SOURCE - À VERIFIER)": "À supprimer (vérifier échantillon avant purge)",
    "À SUPPRIMER (DIFFUSION - MASTER DÉJÀ SUR S3)": "À supprimer",
    "À CHERCHER": "À chercher",
    "S3_KEY À CONSTRUIRE": "S3 key à construire",
    "DOUBLON - À VOIR": "À examiner (doublon)",
}


COLONNES = ["CAS", "UUID", "STATUT", "CHEMIN", "PROBLEMES", "À FAIRE"]

# Noms courts pour la colonne CAS, plus parlants que "CAS 1/2/3" (cf. définition
# des cas dans stagemel_cas_merge.py : présence croisée uuid REF / clé DAO).
CAS_APPARIE = "APPARIÉ"    # cas1 : dans REF et DAO
CAS_DAO_SEUL = "DAO SEUL"  # cas2 : dans DAO seulement
CAS_REF_SEUL = "REF SEUL"  # cas3 : dans REF seulement


def norm_key(k):
    return re.sub(r"[^A-Z0-9]", "", str(k).upper())


def est_a_supprimer(statut):
    return isinstance(statut, str) and statut.startswith("À SUPPRIMER")


def connus_ref(cas1, cas3):
    """key/uuid/path/conservation_statut de cas1+cas3 (fichiers du REF), scindés en
    deux lots : les fichiers pas-encore-supprimés (prioritaires pour un rapprochement)
    et ceux déjà marqués À SUPPRIMER (recours seulement si rien d'autre ne matche —
    proposer comme correspondance un fichier voué à disparaître serait trompeur)."""
    cols = ["key", "uuid", "path", "conservation_statut"]
    parts = [df_[cols] for df_ in (cas1, cas3) if not df_.empty]
    vide = pd.DataFrame(columns=cols)
    if not parts:
        return vide, vide
    ref = pd.concat(parts, ignore_index=True)
    a_supprimer = ref["conservation_statut"].map(est_a_supprimer)
    return ref[~a_supprimer], ref[a_supprimer]


def charger(corpus_code, date):
    paths = {
        cas: dossier_travail(date) / f"{corpus_code}_{cas}_{date}.csv.gz"
        for cas in ("cas1", "cas2", "cas3")
    }
    if not any(p.exists() for p in paths.values()):
        raise FileNotFoundError(
            f"Aucun fichier cas pour {corpus_code} à la date {date} dans {dossier_travail(date)} "
            "— lancer stagemel_cas_merge.py d'abord."
        )
    # Un corpus vide (ni REF ni DAO) donne un recap vide plutôt qu'une erreur :
    # le pipeline enchaîne tous les corpus et ne doit pas s'arrêter sur l'un d'eux.
    return {
        cas: pd.read_csv(p, low_memory=False) if p.exists() else pd.DataFrame()
        for cas, p in paths.items()
    }


def ref_de_l_extraction(date, avertissements):
    """Ref utilisé par stagemel_extraction_corpus.py pour cette date, à défaut le plus récent."""
    trace = ref_trace_path(date)
    if trace.exists():
        # la trace peut venir d'un lancement sous Windows (séparateurs \)
        return trace.read_text(encoding="utf-8").strip().replace("\\", "/")
    avertissements.append(f"{trace} introuvable : masters cherchés dans le ref le plus récent")
    return dernier_ref()


def cles_tif(ref_corpus):
    """Clés (MAJ) des .tif/.tiff parmi les fichiers REF du corpus."""
    est_tif = ref_corpus["extension"].str.lower().isin([".tif", ".tiff"])
    return set(ref_corpus.loc[est_tif, "key"].str.upper())


def jpg_sans_tif(df, tif_keys):
    """.jpg/.jpeg n'ayant pas de .tif/.tiff de même clé dans le corpus."""
    est_jpg = df["extension"].str.lower().isin([".jpg", ".jpeg"])
    return est_jpg & ~df["key"].str.upper().isin(tif_keys)


def a_supprimer_deja_master(df, masters):
    """Lignes À SUPPRIMER (*) dont le master .tif est déjà TRANSFERT_S3_OK."""
    a_master = df["name"].str.rsplit(".", n=1).str[0].str.upper().isin(masters)
    return df["conservation_statut"].map(est_a_supprimer) & a_master


def masters_tif_stems(ref_path):
    """Stems (nom sans extension, MAJ) des .tif déjà TRANSFERT_S3_OK, tout le ref
    confondu (tous corpus) — un À SUPPRIMER dont le master est déjà en sécurité sur
    S3 n'a pas besoin d'être revu, même si ce master est dans un autre corpus_code
    (cf. mémoire project_stagemel_recap : cas MED_MS/AMR_PR, volontaire, pas un bug)."""
    ref = pd.read_csv(ref_path, usecols=["name", "extension", "conservation_statut"], low_memory=False)
    tif_ok = ref["extension"].str.lower().isin([".tif", ".tiff"]) & (ref["conservation_statut"] == "TRANSFERT_S3_OK")
    return set(ref.loc[tif_ok, "name"].str.rsplit(".", n=1).str[0].str.upper())


def build_cas1(df, corpus_code, masters, tif_keys, avertissements):
    if df.empty:
        return pd.DataFrame(columns=["CAS", corpus_code] + COLONNES[1:])

    sans_tif = jpg_sans_tif(df, tif_keys)

    out = pd.DataFrame({
        "CAS": CAS_APPARIE,
        corpus_code: df["name"],
        "UUID": df["uuid"],
        "STATUT": df["conservation_statut"],
        "CHEMIN": df["path"],
        "PROBLEMES": sans_tif.map({True: "Pas de format tif", False: ""}),
        "À FAIRE": df["conservation_statut"].map(ACTION_SUR).fillna(""),
    })

    a_valider = sorted(set(df["conservation_statut"].dropna()) - set(ACTION_SUR))
    for s in a_valider:
        avertissements.append(f"cas1 : action à valider pour le statut « {s} » (politique non définie)")

    # À SUPPRIMER dont le master est déjà sur S3 : rien à décider, on n'encombre pas le recap.
    exclure = a_supprimer_deja_master(df, masters)
    if exclure.any():
        avertissements.append(
            f"cas1 : {int(exclure.sum())} ligne(s) À SUPPRIMER exclue(s) du recap (master déjà TRANSFERT_S3_OK)"
        )
    return out[~exclure.to_numpy()]


def build_cas2(df, corpus_code, cas1, cas3, avertissements):
    if df.empty:
        return pd.DataFrame(columns=["CAS", corpus_code] + COLONNES[1:])

    # Fichiers connus du REF (cas1+cas3), scindés prioritaire / À SUPPRIMER (cf.
    # connus_ref) — un rapprochement ne doit pointer vers un
    # fichier déjà marqué À SUPPRIMER que si aucun autre candidat n'existe.
    prioritaire, secours = connus_ref(cas1, cas3)
    uuid_prioritaire = {norm_key(k): u for k, u in zip(prioritaire["key"], prioritaire["uuid"])}
    uuid_secours = {norm_key(k): u for k, u in zip(secours["key"], secours["uuid"])}

    key_norm = df["key"].map(norm_key)
    dans_prioritaire = key_norm.isin(uuid_prioritaire)
    dans_secours = key_norm.isin(uuid_secours) & ~dans_prioritaire
    candidat = dans_prioritaire | dans_secours
    n_candidats, n_candidats_secours = int(candidat.sum()), int(dans_secours.sum())
    if n_candidats:
        msg = f"cas2 : {n_candidats} candidat(s) ERREUR DAO (clé proche d'un cas1/cas3) — à confirmer via la notice EAD"
        if n_candidats_secours:
            msg += f" (dont {n_candidats_secours} vers un fichier déjà À SUPPRIMER — vérifier en priorité)"
        avertissements.append(msg)

    statut = pd.Series("SEUL DAO", index=df.index)
    statut = statut.mask(dans_prioritaire, "ERREUR DAO (candidat)")
    statut = statut.mask(dans_secours, "ERREUR DAO (candidat, fichier À SUPPRIMER)")
    uuid_col = key_norm.map(uuid_prioritaire)
    uuid_col = uuid_col.where(uuid_col.notna(), key_norm.map(uuid_secours))

    problemes = candidat.map({True: "Clé proche d'un fichier connu (casse/ponctuation ?)", False: "Fichier non retrouvé dans REF"})

    out = pd.DataFrame({
        "CAS": CAS_DAO_SEUL,
        corpus_code: df["nom_fichier_base"],
        "UUID": uuid_col,
        "STATUT": statut,
        "CHEMIN": "",
        "PROBLEMES": problemes,
        "À FAIRE": "",
    })
    return out


def build_cas3(df, corpus_code, masters, tif_keys, avertissements):
    if df.empty:
        return pd.DataFrame(columns=["CAS", corpus_code] + COLONNES[1:])

    problemes = jpg_sans_tif(df, tif_keys).map({
        True: "Absent des notices EAD ; Pas de format tif",
        False: "Absent des notices EAD",
    })
    # À FAIRE volontairement vide : notice à créer ou fichier hors périmètre,
    # c'est à l'utilisateur de trancher.
    out = pd.DataFrame({
        "CAS": CAS_REF_SEUL,
        corpus_code: df["name"],
        "UUID": df["uuid"],
        "STATUT": df["conservation_statut"],
        "CHEMIN": df["path"],
        "PROBLEMES": problemes,
        "À FAIRE": "",
    })

    exclure = a_supprimer_deja_master(df, masters)
    if exclure.any():
        avertissements.append(
            f"cas3 : {int(exclure.sum())} ligne(s) À SUPPRIMER exclue(s) du recap (master déjà TRANSFERT_S3_OK)"
        )
    return out[~exclure.to_numpy()]


FOND_BLANC = PatternFill(fill_type="solid", fgColor="FFFFFFFF")


def appliquer_fond_blanc(ws):
    """Fond blanc systématique (au lieu du fond transparent par défaut d'Excel/openpyxl)."""
    for row in ws.iter_rows():
        for cell in row:
            cell.fill = FOND_BLANC


def ajuster_largeurs(ws, df):
    """Largeur de chaque colonne = texte le plus long (en-tête compris), plafonnée à la limite Excel."""
    for i, col in enumerate(df.columns, start=1):
        longueur = max([len(str(col))] + df[col].dropna().astype(str).str.len().tolist())
        ws.column_dimensions[get_column_letter(i)].width = min(longueur + 2, 255)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("corpus_code")
    parser.add_argument("--date", default=datetime.now().strftime("%Y%m%d"))
    args = parser.parse_args()

    dfs = charger(args.corpus_code, args.date)
    avertissements = []
    masters = masters_tif_stems(ref_de_l_extraction(args.date, avertissements))
    ref_corpus = [d for d in (dfs["cas1"], dfs["cas3"]) if not d.empty]
    tif_keys = cles_tif(pd.concat(ref_corpus)) if ref_corpus else set()
    if all(d.empty for d in dfs.values()):
        avertissements.append(f"{args.corpus_code} : aucun fichier REF ni DAO — recap vide")

    cas1 = build_cas1(dfs["cas1"], args.corpus_code, masters, tif_keys, avertissements)
    cas2 = build_cas2(dfs["cas2"], args.corpus_code, dfs["cas1"], dfs["cas3"], avertissements)
    cas3 = build_cas3(dfs["cas3"], args.corpus_code, masters, tif_keys, avertissements)

    recap = pd.concat([cas1, cas2, cas3], ignore_index=True)
    pivot = recap.groupby(["CAS", "STATUT"], dropna=False).size().rename("nombre").reset_index()

    out_path = dossier_date(args.date) / f"{args.corpus_code.lower()}_recap_{args.date}.xlsx"
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        for nom, df in (("recap", recap), ("pivot", pivot)):
            df.to_excel(writer, sheet_name=nom, index=False)
            appliquer_fond_blanc(writer.sheets[nom])
            ajuster_largeurs(writer.sheets[nom], df)
            writer.sheets[nom].auto_filter.ref = writer.sheets[nom].dimensions  # filtres actifs sur les en-têtes
            writer.sheets[nom].freeze_panes = "A2"  # ligne d'en-tête figée

    print(f"Recap écrit : {out_path}  ({len(recap)} lignes)")
    print(pivot.to_string(index=False))
    if avertissements:
        print("\nÀ valider avant de considérer ce recap comme définitif :")
        for a in avertissements:
            print(f"  - {a}")


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
r"""
stagemel_synthese.py — Synthèse de tous les recaps stagemel d'une date.
-> results/corpus/stagemel/<date>/00_synthese_<date>.xlsx
-> results/corpus/stagemel/<date>/00_synthese_<date>.docx (onglet statuts, un tableau par corpus)

Relit chaque <corpus>_recap_<date>.xlsx (onglet recap) produit par
stagemel_recap.py et en tire quatre onglets :

  - vue_ensemble : une ligne par corpus (+ TOTAL) — fichiers par CAS, candidats
          ERREUR DAO, INCONNU*, fichiers sans À FAIRE, problèmes signalés
          (Pas de format tif, Absent des notices EAD).
  - a_faire : corpus × À FAIRE (« (à trancher) » = cellule vide du recap).
  - statuts : corpus × CAS × STATUT × À FAIRE, en format long (filtrable).
  - a_trancher : fichiers sans À FAIRE regroupées par corpus / CAS / STATUT /
          PROBLEMES — ce qui reste à décider, sans rien trancher ici (cf.
          stagemel_recap.py : la politique INCONNU* n'est jamais inventée).

Usage
-----
  conda run -n rbx-bnr-data python scripts/corpus/stagemel_synthese.py
  conda run -n rbx-bnr-data python scripts/corpus/stagemel_synthese.py --date 20260924
"""
import argparse
import re
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

sys.path.insert(0, str(Path(__file__).parent))
from stagemel_extraction_corpus import dossier_date  # noqa: E402
from stagemel_recap import ajuster_largeurs, appliquer_fond_blanc  # noqa: E402

CAS_ORDRE = ["APPARIÉ", "DAO SEUL", "REF SEUL"]
A_TRANCHER = "(à trancher)"


def charger_recaps(date):
    recaps = {}
    for p in sorted(dossier_date(date).glob(f"*_recap_{date}.xlsx")):
        corpus = re.sub(rf"_recap_{date}\.xlsx$", "", p.name).upper()
        df = pd.read_excel(p, sheet_name="recap", dtype=str).fillna("")
        recaps[corpus] = df
    return recaps


def vue_ensemble(recaps):
    lignes = []
    for corpus, df in recaps.items():
        cas = df["CAS"].value_counts()
        pb = df["PROBLEMES"]
        lignes.append({
            "corpus": corpus,
            "fichiers": len(df),
            **{c: int(cas.get(c, 0)) for c in CAS_ORDRE},
            "ERREUR DAO (candidats)": int(df["STATUT"].str.startswith("ERREUR DAO").sum()),
            "INCONNU*": int(df["STATUT"].str.startswith("INCONNU").sum()),
            "fichiers sans À FAIRE": int((df["À FAIRE"] == "").sum()),
            "Pas de format tif": int(pb.str.contains("Pas de format tif", regex=False).sum()),
            "Absent des notices EAD": int(pb.str.contains("Absent des notices EAD", regex=False).sum()),
        })
    vue = pd.DataFrame(lignes)
    total = vue.drop(columns="corpus").sum().to_frame().T
    total.insert(0, "corpus", "TOTAL")
    return pd.concat([vue, total], ignore_index=True)


def tout(recaps):
    df = pd.concat([d.assign(corpus=c) for c, d in recaps.items()], ignore_index=True)
    df["À FAIRE"] = df["À FAIRE"].replace("", A_TRANCHER)
    return df


def a_faire(df):
    ct = pd.crosstab(df["corpus"], df["À FAIRE"], margins=True, margins_name="TOTAL")
    # colonnes par volume décroissant, TOTAL en dernier
    cols = ct.drop(columns="TOTAL").loc["TOTAL"].sort_values(ascending=False).index.tolist()
    return ct[cols + ["TOTAL"]].reset_index()


def statuts(df):
    return (
        df.groupby(["corpus", "CAS", "STATUT", "À FAIRE"]).size().rename("fichiers").reset_index()
    )


def a_trancher(df):
    reste = df[df["À FAIRE"] == A_TRANCHER]
    return (
        reste.groupby(["corpus", "CAS", "STATUT", "PROBLEMES"]).size().rename("fichiers").reset_index()
        .sort_values(["corpus", "fichiers"], ascending=[True, False])
    )


def nombre_fr(n):
    """1234567 -> '1 234 567' (espace fine insécable)."""
    return f"{int(n):,}".replace(",", "\u202f")


def griser(cell):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "E7E6E6")
    cell._tc.get_or_add_tcPr().append(shd)


def tableau(doc, entetes, lignes, largeurs_cm, droite=(), gras_derniere=False, garder_ensemble=False):
    """Tableau à en-tête grisé répété sur chaque page ; colonnes `droite` alignées à droite."""
    table = doc.add_table(rows=1, cols=len(entetes))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    repete = OxmlElement("w:tblHeader")
    repete.set(qn("w:val"), "true")
    tr_pr.append(repete)
    for i, titre in enumerate(entetes):
        cell = table.rows[0].cells[i]
        cell.text = ""
        cell.paragraphs[0].add_run(titre).bold = True
        griser(cell)
    for n, valeurs in enumerate(lignes):
        cells = table.add_row().cells
        for i, v in enumerate(valeurs):
            run = cells[i].paragraphs[0].add_run(str(v))
            run.bold = gras_derniere and n == len(lignes) - 1
    # largeurs aussi dans la grille du tableau (sinon LibreOffice répartit à parts égales)
    for col, largeur in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), largeurs_cm):
        col.set(qn("w:w"), str(int(Cm(largeur).twips)))
    for r, row in enumerate(table.rows):
        for i, cell in enumerate(row.cells):
            cell.width = Cm(largeurs_cm[i])
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            # tableau d'un seul tenant : chaque ligne reste avec la suivante
            par.paragraph_format.keep_with_next = garder_ensemble and r < len(table.rows) - 1
            if i in droite:
                par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return table


def ecrire_docx(out_path, date, vue, statuts_df):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    for cote in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, cote, Cm(2))
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    date_fr = f"{date[6:]}/{date[4:6]}/{date[:4]}"
    doc.add_heading(f"Synthèse des récapitulatifs stagemel — {date_fr}", level=0)
    total = vue[vue["corpus"] == "TOTAL"].iloc[0]
    doc.add_paragraph(
        f"{len(vue) - 1} récapitulatifs de corpus, {nombre_fr(total['fichiers'])} fichiers. Source\u00a0: "
        f"results/corpus/stagemel/{date}/*_recap_{date}.xlsx (stagemel_recap.py)."
    )
    doc.add_paragraph(
        "Le REF (référentiel des fichiers) est la liste de tous les fichiers de la BnR, "
        "où qu'ils se trouvent\u00a0: S3, Azraël, disques durs, etc. Chaque fichier est "
        "confronté aux fichiers (DAO) cités par les notices EAD, ce qui donne trois cas\u00a0:"
    )
    tableau(doc, ["CAS", "dans le REF", "dans une notice EAD", "signification"], [
        ["APPARIÉ", "oui", "oui", "fichier présent et décrit"],
        ["DAO SEUL", "non", "oui", "fichier cité par une notice mais introuvable dans le REF"],
        ["REF SEUL", "oui", "non", "fichier présent mais décrit par aucune notice"],
    ], [2.6, 2.4, 3.4, 8.6], garder_ensemble=True)
    doc.add_paragraph(
        f"« {A_TRANCHER} » = À FAIRE vide dans le récapitulatif, décision non prise."
    ).paragraph_format.space_before = Pt(6)

    doc.add_heading("Vue d'ensemble", level=1)
    lignes = [
        [r["corpus"], *(nombre_fr(r[c]) for c in ["fichiers", *CAS_ORDRE])]
        for _, r in vue.iterrows()
    ]
    tableau(doc, ["corpus", "fichiers", *CAS_ORDRE], lignes, [4, 3, 3, 3, 3],
            droite=(1, 2, 3, 4), gras_derniere=True)

    ordre_cas = {c: i for i, c in enumerate(CAS_ORDRE)}
    for corpus, d in statuts_df.groupby("corpus", sort=True):
        d = d.assign(_o=d["CAS"].map(ordre_cas)).sort_values(["_o", "fichiers"], ascending=[True, False])
        doc.add_heading(f"{corpus} — {nombre_fr(d['fichiers'].sum())} fichiers", level=1)
        lignes, precedent = [], None
        for _, r in d.iterrows():
            lignes.append([r["CAS"] if r["CAS"] != precedent else "", r["STATUT"], r["À FAIRE"], nombre_fr(r["fichiers"])])
            precedent = r["CAS"]
        lignes.append(["Total", "", "", nombre_fr(d["fichiers"].sum())])
        tableau(doc, ["CAS", "STATUT", "À FAIRE", "fichiers"], lignes, [2.6, 7.2, 5.2, 2],
                droite=(3,), gras_derniere=True, garder_ensemble=True)

    doc.save(out_path)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--date", default=datetime.now().strftime("%Y%m%d"))
    args = parser.parse_args()

    recaps = charger_recaps(args.date)
    if not recaps:
        raise SystemExit(f"Aucun récapitulatif dans {dossier_date(args.date)} — lancer stagemel_recap.py d'abord.")

    df = tout(recaps)
    onglets = {
        "vue_ensemble": vue_ensemble(recaps),
        "a_faire": a_faire(df),
        "statuts": statuts(df),
        "a_trancher": a_trancher(df),
    }

    out_path = dossier_date(args.date) / f"00_synthese_{args.date}.xlsx"
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        for nom, d in onglets.items():
            d.to_excel(writer, sheet_name=nom, index=False)
            ws = writer.sheets[nom]
            appliquer_fond_blanc(ws)
            ajuster_largeurs(ws, d)
            ws.auto_filter.ref = ws.dimensions
            ws.freeze_panes = "B2" if nom in ("vue_ensemble", "a_faire") else "A2"

    docx_path = out_path.with_suffix(".docx")
    ecrire_docx(docx_path, args.date, onglets["vue_ensemble"], onglets["statuts"])

    print(f"Synthèse écrite : {out_path} + {docx_path.name}  ({len(recaps)} récapitulatifs, {len(df)} fichiers)")
    print(onglets["vue_ensemble"].to_string(index=False))


if __name__ == "__main__":
    main()
